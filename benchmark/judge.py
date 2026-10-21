#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
judge.py - LLM-as-judge para a CORRECAO FACTUAL da geracao (complementa Token F1/ROUGE-L).

Le as respostas do sistema (resultados.json, produzido pelo evaluate.py) e as respostas
de referencia (gold standard .xlsx), e pede a um modelo JUIZ que classifique cada par:
    CORRETO   - transmite a mesma informacao essencial (mesmo por outras palavras)
    PARCIAL   - parcialmente correto / incompleto (ex.: lista a que faltam elementos)
    INCORRETO - errado, ausente, ou recusou indevidamente

O juiz deve ser um modelo MAIOR do que o gerador (evita avaliacao circular).
Avalia so as categorias com resposta de referencia (factual, listagem, restricao,
multi_hop, conversacional).

Dependencias: openpyxl + stdlib (urllib). Sem 'requests'.

Variaveis de ambiente:
    GROQ_API_KEY   (obrigatoria)
    JUDGE_MODEL    (opcional; default abaixo - ajusta ao modelo que tens no Groq)

Uso:
    export GROQ_API_KEY=...
    python judge.py --gold gold_standard_corrigido.xlsx --resultados resultados.json
    python judge.py --selftest          # valida a logica sem chamar a API
"""

import os, re, json, time, argparse, urllib.request, urllib.error
from collections import defaultdict
try:
    from dotenv import load_dotenv
    load_dotenv(os.path.join(os.path.dirname(__file__), "..", ".env"))
except ImportError:
    pass

GROQ_URL    = "https://api.groq.com/openai/v1/chat/completions"
JUDGE_MODEL = os.environ.get("JUDGE_MODEL", "openai/gpt-oss-120b")  # modelo de raciocinio (juiz)
API_KEY     = os.environ.get("GROQ_API_KEY", "")

JUDGE_CATS  = {"factual", "listagem", "restricao", "multi_hop", "conversacional"}

PROMPT = """És um avaliador rigoroso de um sistema de perguntas e respostas sobre documentação institucional.
Dada uma PERGUNTA, a RESPOSTA DE REFERÊNCIA (verdade) e a RESPOSTA DO SISTEMA, classifica a correção factual da resposta do sistema.

Critérios:
- CORRETO: transmite a mesma informação essencial da referência, mesmo que por outras palavras ou de forma mais concisa.
- PARCIAL: parcialmente correto ou incompleto (por exemplo, uma lista a que faltam elementos, ou parte da resposta certa).
- INCORRETO: informação errada, em falta, ou o sistema recusou responder quando a referência tem resposta.

Apresenta a classificação final como UMA palavra, no fim da resposta: CORRETO, PARCIAL ou INCORRETO.

PERGUNTA: {q}
RESPOSTA DE REFERÊNCIA: {ref}
RESPOSTA DO SISTEMA: {sys}"""

def classify(txt):
    """Apanha a ULTIMA etiqueta exata no texto (o raciocinio pode conter etiquetas antes)."""
    txt = (txt or "").upper()
    found = None
    for m in re.finditer(r"\b(INCORRETO|PARCIAL|CORRETO)\b", txt):
        found = m.group(1)
    return found or "INCORRETO"

def judge_call(q, ref, sysans, retries=4):
    body = json.dumps({
        "model": JUDGE_MODEL,
        "messages": [{"role": "user", "content": PROMPT.format(q=q, ref=ref, sys=sysans)}],
        "temperature": 0,
        "max_tokens": 1024,   # espaco para o raciocinio + etiqueta final
    }).encode("utf-8")
    req = urllib.request.Request(GROQ_URL, data=body, headers={
        "Authorization": f"Bearer {API_KEY}",
        "Content-Type": "application/json",
        "User-Agent": "python-httpx/0.27.0",
    })
    for attempt in range(retries):
        try:
            with urllib.request.urlopen(req, timeout=120) as r:
                data = json.load(r)
            msg = data["choices"][0]["message"]
            # alguns modelos de raciocinio devolvem o raciocinio num campo separado:
            texto = (msg.get("content") or "") + " " + (msg.get("reasoning") or "")
            return classify(texto)
        except urllib.error.HTTPError as e:
            if e.code == 429 and attempt < retries - 1:
                time.sleep(20 + 15 * attempt); continue  # 20s, 35s, 50s
            raise
    return "INCORRETO"

SCORE = {"CORRETO": 1.0, "PARCIAL": 0.5, "INCORRETO": 0.0}

def load_gold(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True)
    ws = wb["Perguntas"] if "Perguntas" in wb.sheetnames else wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    H = [str(h) for h in rows[0]]
    out = {}
    for r in rows[1:]:
        if not any(r): continue
        d = {H[i]: r[i] for i in range(len(H))}
        out[d["id"]] = d
    return out

def run(gold_path, resultados_path, limit=None, ids=None, out_json="correcao.json"):
    gold = load_gold(gold_path)
    res = json.load(open(resultados_path, encoding="utf-8"))["linhas"]
    alvo = [r for r in res if r.get("categoria") in JUDGE_CATS and r.get("resposta") is not None]
    if ids:
        want = {i.strip() for i in ids.split(",")}
        alvo = [r for r in alvo if r["id"] in want]
    if limit:
        seen = defaultdict(int); kept = []
        for r in alvo:
            if seen[r["categoria"]] < limit:
                kept.append(r); seen[r["categoria"]] += 1
        alvo = kept

    julg = []
    for i, r in enumerate(alvo, 1):
        g = gold.get(r["id"])
        if not g:
            continue
        cls = judge_call(g["pergunta"], g["resposta_esperada"], r["resposta"])
        print(f"[{i}/{len(alvo)}] {r['id']}: {cls}", flush=True)
        julg.append({"id": r["id"], "categoria": r["categoria"], "classe": cls})
        time.sleep(0.3)  # gentil com o rate limit

    # funde com julgamentos anteriores (substitui por id)
    if ids:
        try:
            prev = json.load(open(out_json, encoding="utf-8")).get("linhas", [])
            by_id = {j["id"]: j for j in prev}
            for j in julg: by_id[j["id"]] = j
            julg = list(by_id.values())
            print(f"[merge] fundido com {out_json}: total {len(julg)} linhas", flush=True)
        except FileNotFoundError:
            pass

    # agregacao
    by = defaultdict(list)
    for j in julg: by[j["categoria"]].append(j["classe"])
    agg = {}
    for cat, classes in by.items():
        n = len(classes)
        agg[cat] = {
            "n": n,
            "correto":   sum(c == "CORRETO"   for c in classes) / n,
            "parcial":   sum(c == "PARCIAL"   for c in classes) / n,
            "incorreto": sum(c == "INCORRETO" for c in classes) / n,
            "score":     sum(SCORE[c] for c in classes) / n,
        }
    print("\n==== CORRECAO FACTUAL ====")
    print(json.dumps(agg, indent=2, ensure_ascii=False))
    print("\n==== LATEX ====")
    print(to_latex(agg))
    json.dump({"linhas": julg, "agregado": agg},
              open(out_json, "w", encoding="utf-8"), ensure_ascii=False, indent=2)
    print(f"\n(guardado em {out_json})")

def fmt(x): return f"{x:.2f}"

def to_latex(agg):
    order = ["factual", "listagem", "restricao", "multi_hop", "conversacional"]
    L = ["\\begin{table}[htbp]\\centering\\renewcommand{\\arraystretch}{1.3}",
         "\\rowcolors{2}{ipbrow}{white}",
         "\\begin{tabularx}{\\textwidth}{l *{4}{>{\\centering\\arraybackslash}X}}",
         "\\rowcolor{ipbvinho}",
         "\\textcolor{white}{Categoria} & \\textcolor{white}{Correto} & \\textcolor{white}{Parcial} & "
         "\\textcolor{white}{Incorreto} & \\textcolor{white}{Correção}\\\\"]
    for c in order:
        if c in agg:
            a = agg[c]
            nome = c.replace("_", "-")
            L.append(f"{nome} & {fmt(a['correto'])} & {fmt(a['parcial'])} & {fmt(a['incorreto'])} & {fmt(a['score'])}\\\\")
    L += ["\\bottomrule\\end{tabularx}",
          "\\caption{Correção factual da geração, avaliada por um modelo juiz "
          "(Correção = Correto + 0,5\\,$\\times$\\,Parcial).}",
          "\\label{tab:correcao-geracao}", "\\end{table}"]
    return "\n".join(L)

def selftest():
    global judge_call
    judge_call = lambda q, ref, s, retries=4: "CORRETO" if "15/09" in s or ref.split()[0] in s else "INCORRETO"
    fake_res = {"linhas": [
        {"id": "fact_11", "categoria": "factual", "resposta": "As aulas comecam a 15/09/2025."},
        {"id": "restr_01", "categoria": "restricao", "resposta": "nao sei"},
    ]}
    json.dump(fake_res, open("_res_test.json", "w", encoding="utf-8"), ensure_ascii=False)
    # gold minimo em memoria
    global load_gold
    load_gold = lambda p: {
        "fact_11": {"pergunta": "Quando comecam as aulas?", "resposta_esperada": "As aulas comecam a 15/09/2025."},
        "restr_01": {"pergunta": "Quem atende a segunda?", "resposta_esperada": "Edite Martins Cordeiro."},
    }
    run("_x", "_res_test.json")
    print("\nSELF-TEST OK")

if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--gold", default="gold_standard_corrigido.xlsx")
    ap.add_argument("--resultados", default="resultados.json")
    ap.add_argument("--limit", type=int, default=None)
    ap.add_argument("--ids", default=None, help="re-julga so estes ids (separados por virgula) e funde no correcao.json")
    ap.add_argument("--out", default="correcao.json", help="ficheiro de correcao a ler/gravar")
    ap.add_argument("--selftest", action="store_true")
    args = ap.parse_args()
    if args.selftest:
        selftest()
    elif not API_KEY:
        print("ERRO: define GROQ_API_KEY no ambiente.")
    else:
        run(args.gold, args.resultados, limit=args.limit, ids=args.ids, out_json=args.out)
