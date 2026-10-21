#!/usr/bin/env python3
"""
benchmark/run_benchmark.py — Avaliação automática do EducaAssistAI (LLM-as-a-Judge)

Interliga com o EducaAssistAI exclusivamente via HTTP — sem imports do projecto.

Uso:
    python benchmark/run_benchmark.py --input perguntas.xlsx

Opções:
    --input        Ficheiro Excel (.xlsx) com colunas: Pergunta, Resposta Esperada, Fonte
    --agent-url    URL base do EducaAssistAI (default: http://localhost:8000)
    --judge        LLM avaliador: groq | gemini  (default: groq)
    --judge-model  Modelo do avaliador (default do provider se omitido)
    --output-dir   Directório de saída (default: benchmark/resultados)
    --timeout      Timeout por pedido ao agente em segundos (default: 60)
    --retries      Nº de tentativas por pedido (default: 3)
    --delay        Atraso entre perguntas em segundos (default: 1.0)
"""

from __future__ import annotations

import argparse
import json
import logging
import os
import sys
import time
from datetime import datetime
from pathlib import Path
from typing import Optional

import openpyxl
import requests
from dotenv import load_dotenv

# ─── Logging ──────────────────────────────────────────────────────────────────

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s  %(message)s",
    datefmt="%H:%M:%S",
)
log = logging.getLogger("benchmark")

# ─── .env (raiz do projecto, um nível acima de benchmark/) ────────────────────

_env_path = Path(__file__).parent.parent / ".env"
if _env_path.exists():
    load_dotenv(_env_path)

# ─────────────────────────────────────────────────────────────────────────────
# 1. LEITURA DO EXCEL
# ─────────────────────────────────────────────────────────────────────────────

def load_questions(path: str) -> list[dict]:
    """Lê o ficheiro Excel e devolve lista de dicts com keys:
    pergunta, resposta_esperada, fonte.

    Reconhece cabeçalhos com nomes flexíveis (case-insensitive) e salta
    linhas completamente vazias.
    """
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    if not rows:
        raise ValueError("Ficheiro Excel está vazio.")

    # Detectar cabeçalho (primeira linha)
    header = [str(c).strip().lower() if c is not None else "" for c in rows[0]]
    col_map: dict[str, int] = {}
    for i, h in enumerate(header):
        if "pergunta" in h or "question" in h:
            col_map.setdefault("pergunta", i)
        elif ("resposta" in h and "esperada" in h) or "expected" in h or "ground" in h:
            col_map.setdefault("resposta_esperada", i)
        elif "fonte" in h or "source" in h or "origem" in h:
            col_map.setdefault("fonte", i)

    missing = [k for k in ("pergunta", "resposta_esperada") if k not in col_map]
    if missing:
        raise ValueError(
            f"Colunas em falta no Excel: {missing}.\n"
            f"Cabeçalho detectado: {header}\n"
            f"Nomes aceites — Pergunta: 'pergunta'/'question'; "
            f"Resposta Esperada: 'resposta esperada'/'expected answer'/'ground truth'; "
            f"Fonte: 'fonte'/'source'."
        )

    questions: list[dict] = []
    for row in rows[1:]:
        def _cell(key: str) -> str:
            idx = col_map.get(key)
            if idx is None or idx >= len(row) or row[idx] is None:
                return ""
            return str(row[idx]).strip()

        pergunta = _cell("pergunta")
        if not pergunta or pergunta.lower() == "none":
            continue

        questions.append({
            "pergunta":           pergunta,
            "resposta_esperada":  _cell("resposta_esperada"),
            "fonte":              _cell("fonte"),
        })

    log.info(f"Carregadas {len(questions)} perguntas de '{path}'")
    return questions


# ─────────────────────────────────────────────────────────────────────────────
# 2. PEDIDO AO AGENTE
# ─────────────────────────────────────────────────────────────────────────────

def query_agent(
    question: str,
    base_url: str,
    timeout: int = 60,
    max_retries: int = 3,
) -> tuple[str, list[str]]:
    """POST /query → (answer, sources).

    Backoff exponencial em erros de rede e respostas 5xx.
    Lança RuntimeError se todos os retries falharem.
    """
    url = f"{base_url.rstrip('/')}/query"
    payload = {"question": question}

    for attempt in range(1, max_retries + 1):
        try:
            resp = requests.post(url, json=payload, timeout=timeout)
            if resp.status_code == 200:
                data = resp.json()
                return data.get("answer", ""), data.get("sources", [])
            if resp.status_code < 500:
                # 4xx — erro de cliente, retry não adianta
                raise RuntimeError(
                    f"Agente devolveu {resp.status_code}: {resp.text[:300]}"
                )
            log.warning(f"Agente devolveu {resp.status_code} (tentativa {attempt}/{max_retries})")
        except requests.exceptions.Timeout:
            log.warning(f"Timeout ao contactar o agente (tentativa {attempt}/{max_retries})")
        except requests.exceptions.ConnectionError:
            log.warning(f"Ligação recusada ao agente (tentativa {attempt}/{max_retries})")
        except RuntimeError:
            raise
        except Exception as exc:
            log.warning(f"Erro inesperado no pedido ao agente: {exc} (tentativa {attempt}/{max_retries})")

        if attempt < max_retries:
            wait = 2 ** attempt
            log.info(f"      A aguardar {wait}s antes de nova tentativa…")
            time.sleep(wait)

    raise RuntimeError(
        f"O agente não respondeu após {max_retries} tentativas ('{question[:60]}…')"
    )


# ─────────────────────────────────────────────────────────────────────────────
# 3. LLM AVALIADOR (Judge)
# ─────────────────────────────────────────────────────────────────────────────

_JUDGE_SYSTEM = """\
És um avaliador rigoroso de sistemas de perguntas e respostas (QA).

A tua função é comparar uma RESPOSTA OBTIDA de um sistema de IA com a RESPOSTA ESPERADA \
(ground truth), para uma dada PERGUNTA.

CRITÉRIOS:
1. CORREÇÃO FACTUAL — a informação essencial da resposta esperada está presente na resposta obtida?
2. COBERTURA — a resposta obtida aborda os pontos principais da resposta esperada?
3. NÃO CONTRADIÇÃO — a resposta obtida não contradiz a resposta esperada em factos relevantes?

NOTAS:
- Ignora diferenças de formatação, ordem ou estilo de escrita.
- A resposta obtida pode conter MAIS informação que a esperada — não é penalizado.
- Considera "match: true" se a pergunta é respondida correctamente, mesmo com outras palavras.
- Considera "match: false" se a resposta está errada, vazia, contradiz factos, ou declara \
não ter informação quando a resposta esperada existe.

Responde APENAS com um objecto JSON válido, sem markdown, sem texto adicional:
{
  "match": true ou false,
  "justificacao": "Explicação concisa em 1-2 frases."
}"""

_JUDGE_USER = """\
PERGUNTA: {pergunta}

RESPOSTA ESPERADA: {resposta_esperada}

RESPOSTA OBTIDA: {resposta_obtida}"""


def _parse_judge_response(raw: str) -> dict:
    """Extrai JSON da resposta do avaliador, tolerando markdown fences."""
    raw = raw.strip()
    # remover ```json ... ``` se presente
    if raw.startswith("```"):
        raw = raw.split("```")[1]
        if raw.startswith("json"):
            raw = raw[4:]
    return json.loads(raw.strip())


def judge_with_groq(
    pergunta: str,
    resposta_esperada: str,
    resposta_obtida: str,
    model: str = "llama-3.1-8b-instant",
    max_retries: int = 3,
) -> dict:
    """Avalia via Groq Chat Completions. Requer GROQ_API_KEY no .env."""
    api_key = os.getenv("GROQ_API_KEY", "").strip()
    if not api_key:
        raise RuntimeError("GROQ_API_KEY não definida no .env")

    url = "https://api.groq.com/openai/v1/chat/completions"
    headers = {"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"}
    payload = {
        "model": model,
        "messages": [
            {"role": "system", "content": _JUDGE_SYSTEM},
            {"role": "user", "content": _JUDGE_USER.format(
                pergunta=pergunta,
                resposta_esperada=resposta_esperada,
                resposta_obtida=resposta_obtida,
            )},
        ],
        "temperature": 0,
        "max_tokens": 350,
        "response_format": {"type": "json_object"},
    }

    for attempt in range(1, max_retries + 1):
        try:
            resp = requests.post(url, headers=headers, json=payload, timeout=30)
            if resp.status_code == 200:
                content = resp.json()["choices"][0]["message"]["content"]
                return _parse_judge_response(content)
            if resp.status_code == 429:
                wait = 15 * attempt
                log.warning(f"Rate limit Groq — aguardar {wait}s (tentativa {attempt})")
                time.sleep(wait)
                continue
            log.warning(f"Groq devolveu {resp.status_code}: {resp.text[:200]}")
        except json.JSONDecodeError as exc:
            log.warning(f"Avaliador Groq devolveu JSON inválido: {exc}")
        except Exception as exc:
            log.warning(f"Erro avaliador Groq: {exc} (tentativa {attempt}/{max_retries})")

        if attempt < max_retries:
            time.sleep(2 ** attempt)

    return {"match": False, "justificacao": f"[Erro: avaliador Groq não respondeu após {max_retries} tentativas]"}


def judge_with_gemini(
    pergunta: str,
    resposta_esperada: str,
    resposta_obtida: str,
    model: str = "gemini-2.0-flash",
    max_retries: int = 3,
) -> dict:
    """Avalia via Google Gemini REST. Requer GOOGLE_API_KEY no .env."""
    api_key = os.getenv("GOOGLE_API_KEY", "").strip()
    if not api_key:
        raise RuntimeError("GOOGLE_API_KEY não definida no .env")

    url = f"https://generativelanguage.googleapis.com/v1beta/models/{model}:generateContent"
    params = {"key": api_key}
    combined_prompt = (
        _JUDGE_SYSTEM
        + "\n\n"
        + _JUDGE_USER.format(
            pergunta=pergunta,
            resposta_esperada=resposta_esperada,
            resposta_obtida=resposta_obtida,
        )
    )
    payload = {
        "contents": [{"parts": [{"text": combined_prompt}]}],
        "generationConfig": {
            "temperature": 0,
            "maxOutputTokens": 350,
            "responseMimeType": "application/json",
        },
    }

    for attempt in range(1, max_retries + 1):
        try:
            resp = requests.post(url, params=params, json=payload, timeout=30)
            if resp.status_code == 200:
                content = resp.json()["candidates"][0]["content"]["parts"][0]["text"]
                return _parse_judge_response(content)
            if resp.status_code == 429:
                wait = 15 * attempt
                log.warning(f"Rate limit Gemini — aguardar {wait}s (tentativa {attempt})")
                time.sleep(wait)
                continue
            log.warning(f"Gemini devolveu {resp.status_code}: {resp.text[:200]}")
        except json.JSONDecodeError as exc:
            log.warning(f"Avaliador Gemini devolveu JSON inválido: {exc}")
        except Exception as exc:
            log.warning(f"Erro avaliador Gemini: {exc} (tentativa {attempt}/{max_retries})")

        if attempt < max_retries:
            time.sleep(2 ** attempt)

    return {"match": False, "justificacao": f"[Erro: avaliador Gemini não respondeu após {max_retries} tentativas]"}


# ─────────────────────────────────────────────────────────────────────────────
# 4. EXPORTAÇÃO DE RESULTADOS
# ─────────────────────────────────────────────────────────────────────────────

def export_results(results: list[dict], output_dir: str, timestamp: str) -> tuple[str, str]:
    """Escreve Excel + JSON. Devolve (excel_path, json_path)."""
    from openpyxl.styles import Alignment, Font, PatternFill

    out = Path(output_dir)
    out.mkdir(parents=True, exist_ok=True)

    excel_path = str(out / f"resultados_benchmark_{timestamp}.xlsx")
    json_path  = str(out / f"resultados_benchmark_{timestamp}.json")

    # ── JSON ──────────────────────────────────────────────────────────────────
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(results, f, ensure_ascii=False, indent=2)

    # ── Excel ─────────────────────────────────────────────────────────────────
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Benchmark"

    # Estilos
    h_font  = Font(bold=True, color="FFFFFF", size=11)
    h_fill  = PatternFill("solid", fgColor="1D4ED8")   # azul escuro
    ok_fill = PatternFill("solid", fgColor="D1FAE5")   # verde claro
    ko_fill = PatternFill("solid", fgColor="FEE2E2")   # vermelho claro
    er_fill = PatternFill("solid", fgColor="FEF3C7")   # amarelo claro
    wrap    = Alignment(wrap_text=True, vertical="top")

    COLUMNS = [
        ("Nº",                4),
        ("Pergunta",         42),
        ("Resposta Esperada", 42),
        ("Fonte",            22),
        ("Resposta Obtida",  52),
        ("Match",             9),
        ("Justificação",     46),
        ("Erro Agente",      28),
        ("Duração (s)",      13),
    ]

    for ci, (name, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=ci, value=name)
        cell.font = h_font
        cell.fill = h_fill
        cell.alignment = wrap
        ws.column_dimensions[openpyxl.utils.get_column_letter(ci)].width = width

    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"

    for ri, r in enumerate(results, start=2):
        match_val = r.get("match")
        row_fill  = ok_fill if match_val is True else (ko_fill if match_val is False else er_fill)
        match_str = "✓" if match_val is True else ("✗" if match_val is False else "?")

        values = [
            ri - 1,
            r.get("pergunta", ""),
            r.get("resposta_esperada", ""),
            r.get("fonte", ""),
            r.get("resposta_obtida", ""),
            match_str,
            r.get("justificacao", ""),
            r.get("erro_agente", ""),
            r.get("duracao_s", ""),
        ]

        for ci, val in enumerate(values, start=1):
            cell = ws.cell(row=ri, column=ci, value=val)
            cell.alignment = wrap
            # cor na coluna Match
            if ci == 6:
                cell.fill = row_fill
                cell.font = Font(bold=True)

        ws.row_dimensions[ri].height = 55

    # Linha de resumo
    total   = len(results)
    correct = sum(1 for r in results if r.get("match") is True)
    errs    = sum(1 for r in results if r.get("erro_agente"))
    evaluated = sum(1 for r in results if r.get("match") is not None)

    sr = total + 3
    ws.cell(row=sr, column=1, value="RESUMO").font = Font(bold=True, size=11)
    ws.cell(row=sr, column=2, value=(
        f"Total: {total}  |  Avaliadas: {evaluated}  |  "
        f"Corretas: {correct}  |  Incorretas: {evaluated - correct}  |  "
        f"Erros agente: {errs}  |  "
        f"Precisão: {correct / evaluated * 100:.1f}%" if evaluated else "Precisão: N/A"
    )).font = Font(bold=True)

    wb.save(excel_path)

    return excel_path, json_path


# ─────────────────────────────────────────────────────────────────────────────
# 5. MAIN
# ─────────────────────────────────────────────────────────────────────────────

def main() -> None:
    parser = argparse.ArgumentParser(
        description="Benchmark automático do EducaAssistAI via LLM-as-a-Judge",
        formatter_class=argparse.ArgumentDefaultsHelpFormatter,
    )
    parser.add_argument("--input",        required=True,              help="Ficheiro Excel de entrada (.xlsx)")
    parser.add_argument("--agent-url",    default="http://localhost:8000", help="URL base do agente")
    parser.add_argument("--judge",        default="groq", choices=["groq", "gemini"], help="LLM avaliador")
    parser.add_argument("--judge-model",  default="",                 help="Modelo do avaliador (usa default se vazio)")
    parser.add_argument("--output-dir",   default="benchmark/resultados", help="Directório de saída")
    parser.add_argument("--timeout",      type=int,   default=60,     help="Timeout pedido ao agente (s)")
    parser.add_argument("--retries",      type=int,   default=3,      help="Nº de tentativas por pedido")
    parser.add_argument("--delay",        type=float, default=1.0,    help="Atraso entre perguntas (s)")
    args = parser.parse_args()

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")

    # ── 1. Carregar perguntas ─────────────────────────────────────────────────
    try:
        questions = load_questions(args.input)
    except Exception as exc:
        log.error(f"Erro a ler o ficheiro Excel: {exc}")
        sys.exit(1)

    if not questions:
        log.error("O ficheiro Excel não contém perguntas válidas.")
        sys.exit(1)

    total = len(questions)

    # ── 2. Verificar conectividade com o agente ───────────────────────────────
    health_url = f"{args.agent_url.rstrip('/')}/health"
    try:
        r = requests.get(health_url, timeout=10)
        status = "online" if r.status_code == 200 else f"status {r.status_code}"
        log.info(f"Agente {status} em {args.agent_url}")
    except Exception as exc:
        log.error(f"Não foi possível ligar ao agente em {args.agent_url}: {exc}")
        log.error("Certifica-te de que o backend está a correr (uvicorn app:app --reload).")
        sys.exit(1)

    # ── 3. Configurar avaliador ───────────────────────────────────────────────
    if args.judge == "groq":
        judge_model = args.judge_model or "llama-3.1-8b-instant"
        def judge_fn(p: str, re: str, ro: str) -> dict:
            return judge_with_groq(p, re, ro, model=judge_model, max_retries=args.retries)
    else:
        judge_model = args.judge_model or "gemini-2.0-flash"
        def judge_fn(p: str, re: str, ro: str) -> dict:
            return judge_with_gemini(p, re, ro, model=judge_model, max_retries=args.retries)

    log.info(f"Avaliador: {args.judge.upper()} / {judge_model}")
    log.info(f"A iniciar benchmark — {total} pergunta(s)")
    log.info("─" * 65)

    # ── 4. Loop principal ─────────────────────────────────────────────────────
    results: list[dict] = []

    for idx, q in enumerate(questions, start=1):
        pergunta          = q["pergunta"]
        resposta_esperada = q["resposta_esperada"]
        fonte             = q["fonte"]

        prefix = f"[{idx}/{total}]"
        log.info(f"{prefix} {pergunta[:72]}{'…' if len(pergunta) > 72 else ''}")

        # 4a. Consultar o agente
        t0 = time.perf_counter()
        resposta_obtida = ""
        sources: list[str] = []
        erro_agente = ""

        try:
            resposta_obtida, sources = query_agent(
                pergunta,
                base_url=args.agent_url,
                timeout=args.timeout,
                max_retries=args.retries,
            )
            preview = resposta_obtida[:90] + ("…" if len(resposta_obtida) > 90 else "")
            log.info(f"        Agente → {preview}")
        except Exception as exc:
            erro_agente = str(exc)
            log.error(f"        Erro agente: {erro_agente}")

        duracao = round(time.perf_counter() - t0, 2)

        # 4b. Avaliar com LLM
        match: Optional[bool] = None
        justificacao = ""

        if not erro_agente:
            try:
                verdict      = judge_fn(pergunta, resposta_esperada, resposta_obtida)
                match        = bool(verdict.get("match", False))
                justificacao = str(verdict.get("justificacao", ""))
                icon = "✓" if match else "✗"
                log.info(f"        Avaliação {icon} — {justificacao[:80]}")
            except Exception as exc:
                justificacao = f"[Erro avaliador: {exc}]"
                log.error(f"        Erro avaliador: {exc}")

        results.append({
            "pergunta":           pergunta,
            "resposta_esperada":  resposta_esperada,
            "fonte":              fonte,
            "resposta_obtida":    resposta_obtida,
            "sources_agente":     sources,
            "match":              match,
            "justificacao":       justificacao,
            "erro_agente":        erro_agente,
            "duracao_s":          duracao,
        })

        if idx < total and args.delay > 0:
            time.sleep(args.delay)

    # ── 5. Estatísticas finais ────────────────────────────────────────────────
    log.info("─" * 65)
    evaluated = sum(1 for r in results if r["match"] is not None)
    correct   = sum(1 for r in results if r["match"] is True)
    errs      = sum(1 for r in results if r["erro_agente"])

    log.info("RESULTADOS FINAIS:")
    log.info(f"  Total de perguntas  : {total}")
    log.info(f"  Avaliadas           : {evaluated}")
    log.info(f"  Corretas   (✓)      : {correct}")
    log.info(f"  Incorretas (✗)      : {evaluated - correct}")
    log.info(f"  Erros no agente     : {errs}")
    if evaluated:
        log.info(f"  Precisão            : {correct / evaluated * 100:.1f}%")

    # ── 6. Exportar ───────────────────────────────────────────────────────────
    try:
        excel_path, json_path = export_results(results, args.output_dir, timestamp)
        log.info("")
        log.info("Ficheiros guardados:")
        log.info(f"  Excel → {excel_path}")
        log.info(f"  JSON  → {json_path}")
    except Exception as exc:
        log.error(f"Erro a exportar resultados: {exc}")
        # Guardar pelo menos o JSON em raw como fallback
        fallback = Path(args.output_dir) / f"fallback_{timestamp}.json"
        fallback.parent.mkdir(parents=True, exist_ok=True)
        with open(fallback, "w", encoding="utf-8") as f:
            json.dump(results, f, ensure_ascii=False, indent=2)
        log.info(f"  Fallback JSON → {fallback}")


if __name__ == "__main__":
    main()
